import os
import re
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

from openpyxl.chart import ScatterChart, Reference, Series

from PySide6.QtCore import QTimer 

from PySide6.QtWidgets import (
    QMainWindow, QVBoxLayout, QHBoxLayout, QWidget, QPushButton,
    QFileDialog, QMessageBox, QLabel, QListWidget, QTableWidget,
    QTableWidgetItem, QCheckBox, QComboBox, QDialog, QInputDialog
)

from dialogs import GeometryDialog, PreprocessOptionsDialog
from preprocess import preprocess_data
from utils import deform_qform_table, force_displacement_curve_from_raw

class DataOptimizer(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Flow Curve Tool")
        self.setGeometry(200, 200, 1300, 650)

        self.datasets = []
        self.current_index = None

        # ---------- Layout ----------
        main_layout = QHBoxLayout()

        # Preview table (left)
        self.preview_table = QTableWidget()
        self.preview_table.setRowCount(0)
        self.preview_table.setColumnCount(0)
        main_layout.addWidget(self.preview_table, 3)

        # Sidebar (right)
        sidebar = QVBoxLayout()

        self.dataset_list = QListWidget()
        self.dataset_list.currentRowChanged.connect(self.change_dataset)
        self.dataset_list.currentRowChanged.connect(self.update_buttons)
        sidebar.addWidget(QLabel("Loaded Datasets:"))
        sidebar.addWidget(self.dataset_list)

        self.load_btn = QPushButton("Add CSV + Parameters")
        self.load_btn.clicked.connect(self.load_dataset)
        sidebar.addWidget(self.load_btn)

        self.edit_params_button = QPushButton("Edit Specimen Parameters")
        self.edit_params_button.clicked.connect(self.edit_parameters)
        self.edit_params_button.setEnabled(False)
        sidebar.addWidget(self.edit_params_button)

        self.remove_button = QPushButton("Remove Dataset from List")
        self.remove_button.clicked.connect(self.remove_dataset)
        self.remove_button.setEnabled(False)
        sidebar.addWidget(self.remove_button)

        self.clear_all_btn = QPushButton("Clear All Datasets")
        self.clear_all_btn.clicked.connect(self.clear_all_datasets)
        self.clear_all_btn.setEnabled(False)
        sidebar.addWidget(self.clear_all_btn)

        self.preprocess_btn = QPushButton("Process Raw Data (Selected)")
        self.preprocess_btn.clicked.connect(self.run_preprocess)
        self.preprocess_btn.setEnabled(False)
        sidebar.addWidget(self.preprocess_btn)
        
        self.chk_process_all = QCheckBox("Process all loaded datasets")
        self.chk_process_all.setChecked(False)
        sidebar.addWidget(self.chk_process_all)

        # Plot filter checkboxes
        sidebar.addWidget(QLabel("Curves to show in plot:"))
        self.chk_none = QCheckBox("Unfiltered (reference)")
        self.chk_mild = QCheckBox("Mild filter (Savgol)")
        self.chk_strong = QCheckBox("Strong filter (Butterworth)")
        self.chk_none.setChecked(True)
        self.chk_mild.setChecked(True)
        self.chk_strong.setChecked(False)
        sidebar.addWidget(self.chk_none)
        sidebar.addWidget(self.chk_mild)
        sidebar.addWidget(self.chk_strong)

        self.plot_btn = QPushButton("Plots")
        self.plot_btn.clicked.connect(self.show_plots)
        self.plot_btn.setEnabled(False)
        sidebar.addWidget(self.plot_btn)

        # -------- Export --------
        sidebar.addWidget(QLabel("Export filter:"))
        self.export_filter_combo = QComboBox()
        self.export_filter_combo.addItems(["none", "mild", "strong"])
        sidebar.addWidget(self.export_filter_combo)

        self.export_btn = QPushButton("Export (Excel)")
        self.export_btn.clicked.connect(self.export_results)
        self.export_btn.setEnabled(False)
        sidebar.addWidget(self.export_btn)
        
        self.force_disp_export_btn = QPushButton("Export Force–Displacement (Selected)")
        self.force_disp_export_btn.clicked.connect(self.export_force_displacement_selected)
        self.force_disp_export_btn.setEnabled(False)
        sidebar.addWidget(self.force_disp_export_btn)

        self.deform_export_btn = QPushButton("DEFORM / QForm Export")
        self.deform_export_btn.clicked.connect(self.export_deform_qform)
        self.deform_export_btn.setEnabled(False)
        sidebar.addWidget(self.deform_export_btn)
        
        self.extrapolate_btn = QPushButton("Extrapolate Models (Selected)")
        self.extrapolate_btn.clicked.connect(self.run_extrapolation)
        self.extrapolate_btn.setEnabled(False)
        sidebar.addWidget(self.extrapolate_btn)

        self.deform_key_export_btn = QPushButton("DEFORM 3D Keyword Export (.key)")
        self.deform_key_export_btn.clicked.connect(self.export_deform_key)
        self.deform_key_export_btn.setEnabled(False)
        sidebar.addWidget(self.deform_key_export_btn)

        sidebar.addStretch()
        main_layout.addLayout(sidebar, 1)

        container = QWidget()
        container.setLayout(main_layout)
        self.setCentralWidget(container)

    # ---------- Helpers ----------

    def update_list_entry(self, idx: int):
        ds = self.datasets[idx]

        dx = ds.get("strain_increment", None)
        dx_txt = "Δε=?" if dx is None else f"Δε={dx:g}"

        entry = (
            f"{ds['filename']}  |  D={ds['diameter']} mm, H={ds['height']} mm  |  "
            f"{dx_txt}"
        )

        if idx < self.dataset_list.count():
            self.dataset_list.item(idx).setText(entry)
        else:
            self.dataset_list.addItem(entry)

    def show_preview(self):
        if self.current_index is None or self.current_index < 0 or self.current_index >= len(self.datasets):
            self.preview_table.clear()
            self.preview_table.setRowCount(0)
            self.preview_table.setColumnCount(0)
            return

        ds = self.datasets[self.current_index]

        if ds.get("preprocessed", False) and ds.get("proc_data") is not None:
            df = ds["proc_data"]
        else:
            df = ds["raw_data"]

        preview = df.head(20)

        self.preview_table.setRowCount(len(preview))
        self.preview_table.setColumnCount(len(preview.columns))
        self.preview_table.setHorizontalHeaderLabels(preview.columns)

        for i in range(len(preview)):
            for j in range(len(preview.columns)):
                item = QTableWidgetItem(str(preview.iat[i, j]))
                self.preview_table.setItem(i, j, item)

    def refresh_ui_state(self):
        has_selection = self.current_index is not None and 0 <= self.current_index < len(self.datasets)
        any_processed = any(d.get("preprocessed", False) and d.get("proc_data") is not None for d in self.datasets)

        self.clear_all_btn.setEnabled(len(self.datasets) > 0)
        self.preprocess_btn.setEnabled(has_selection)
        
        self.plot_btn.setEnabled(any_processed)
        self.export_btn.setEnabled(any_processed)
        
        has_proc_sel = (
            has_selection   
            and self.datasets[self.current_index].get("preprocessed", False)
            and self.datasets[self.current_index].get("proc_data") is not None
            )
        self.deform_export_btn.setEnabled(has_proc_sel)
        self.deform_key_export_btn.setEnabled(has_proc_sel)
        self.force_disp_export_btn.setEnabled(has_proc_sel)

        has_proc_sel = (
            has_selection   
            and self.datasets[self.current_index].get("preprocessed", False)
            and self.datasets[self.current_index].get("proc_data") is not None
            )
    
        self.extrapolate_btn.setEnabled(has_proc_sel)

    # -------------------- Button Enable/Disable --------------------
    def update_buttons(self, index):
        has_selection = index >= 0 and index < len(self.datasets)
        self.edit_params_button.setEnabled(has_selection)
        self.remove_button.setEnabled(has_selection)

    # -------------------- Load CSV --------------------
    def load_dataset(self):
        paths, _ = QFileDialog.getOpenFileNames(
            self,
            "Open CSV files (multi-select)",
            "",
            "CSV Files (*.csv)"
            )
        if not paths:
            return

        dlg = GeometryDialog(self)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return

        diameter, height = dlg.get_values()
        if diameter is None or height is None:
            QMessageBox.warning(self, "Error", "Please enter valid numbers!")
            return

        loaded = 0
        failed = []

        for path in paths:
            try:
                raw = pd.read_csv(path)
                if raw is None or raw.empty:
                    raise ValueError("CSV is empty.")

                ds = {
                    "filename": os.path.basename(path),
                    "raw_data": raw,
                    "diameter": float(diameter),
                    "height": float(height),
                    "proc_data": None,
                    "preprocessed": False,
                }

                self.datasets.append(ds)
                self.update_list_entry(len(self.datasets) - 1)
                loaded += 1

            except Exception as e:
                failed.append(f"{os.path.basename(path)}: {e}")

        if loaded > 0:
            self.refresh_ui_state()
            QMessageBox.information(self, "Success", f"Loaded {loaded} file(s).")

        if failed:
            QMessageBox.warning(
                self,
                "Some files failed",
                "The following files could not be loaded:\n\n" + "\n".join(failed)
            )

    # -------------------- Switch Dataset --------------------
    def change_dataset(self, idx: int):
        self.current_index = idx
        self.refresh_ui_state()
        self.show_preview()

    # -------------------- Edit / Remove --------------------
    def edit_parameters(self):
        if self.current_index is None:
            return

        # NEW: Check if there are multiple datasets and ask for Probendaten
        if len(self.datasets) > 1:
            msgBox = QMessageBox(self)
            msgBox.setWindowTitle("Edit Specimen Parameters")
            msgBox.setText("You have multiple datasets loaded.\nHow would you like to edit the geometry?")
            
            btn_manual = msgBox.addButton("Edit Selected Manually", QMessageBox.ActionRole)
            btn_auto = msgBox.addButton("Auto-Match All (Probendaten Excel)", QMessageBox.ActionRole)
            btn_cancel = msgBox.addButton("Cancel", QMessageBox.RejectRole)
            
            msgBox.exec()
            
            if msgBox.clickedButton() == btn_cancel:
                return
            elif msgBox.clickedButton() == btn_auto:
                self._import_probendaten()
                return

        # Standard manual edit
        ds = self.datasets[self.current_index]
        dialog = GeometryDialog(self, ds["diameter"], ds["height"])
        if dialog.exec() == QDialog.Accepted:
            diameter, height = dialog.get_values()
            if diameter is None or height is None:
                QMessageBox.warning(self, "Error", "Please enter valid numbers!")
                return
            ds["diameter"] = diameter
            ds["height"] = height
            self.update_list_entry(self.current_index)
            QMessageBox.information(self, "Info", "Parameters updated successfully.")
            if ds["preprocessed"]:
                QMessageBox.warning(self, "Warning", "Raw data already processed. Reprocess if needed.")

    def _import_probendaten(self):
        """Internal worker function to load and match the Probendaten Excel file"""
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Open Probendaten File", "", "Excel Files (*.xlsx *.xls);;CSV Files (*.csv)"
        )
        if not file_path:
            return

        try:
            if file_path.lower().endswith('.csv'):
                df_probe = pd.read_csv(file_path)
            else:
                df_probe = pd.read_excel(file_path)
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Could not read file:\n{str(e)}")
            return

        v_col = None
        d_col = None
        h_col = None

        for c in df_probe.columns:
            c_str = str(c).strip().lower()
            
            # Match "Versuch" or "Versuchsnummer"
            if 'versuch' in c_str:
                v_col = c
            
            # Match Diameter: contains 'ø', 'durchmesser', or 'd0'
            elif 'ø' in c_str or 'durchmesser' in c_str or 'd0' in c_str:
                d_col = c
                
            # Match Height: contains 'h0', 'höh', or 'hoeh'
            elif 'h0' in c_str or 'höh' in c_str or 'hoeh' in c_str:
                h_col = c

        if not v_col or not d_col or not h_col:
            QMessageBox.warning(
                self, "Error", 
                f"Could not find the required columns in the file.\n\n"
                f"Found columns: {', '.join(str(c) for c in df_probe.columns)}\n"
                f"Looking for: 'Versuch', 'Ø' / 'd0' (Diameter), and 'h0' (Height)"
            )
            return

        matched_count = 0
        already_processed = False

        for i, ds in enumerate(self.datasets):
            fname = ds["filename"]
            base_fname = fname.lower().replace('.csv', '')
            
            # Extract prefix number from filename e.g. "015-T100.csv" -> 15
            prefix_match = re.match(r'^0*(\d+)', fname)
            prefix_int = int(prefix_match.group(1)) if prefix_match else None

            for idx, row in df_probe.iterrows():
                v_val = row[v_col]
                if pd.isna(v_val): continue
                v_str = str(v_val).strip()

                is_match = False
                # Method 1: Strict Integer prefix match (Safest)
                if prefix_int is not None:
                    try:
                        if int(float(v_str)) == prefix_int:
                            is_match = True
                    except ValueError:
                        pass
                
                # Method 2: String prefix match (Fallback)
                if not is_match:
                    if base_fname == v_str.lower() or base_fname.startswith(v_str.lower() + "-") or base_fname.startswith(v_str.lower() + "_"):
                        is_match = True

                if is_match:
                    d_val = row[d_col]
                    h_val = row[h_col]
                    if pd.notna(d_val) and pd.notna(h_val):
                        ds["diameter"] = float(d_val)
                        ds["height"] = float(h_val)
                        self.update_list_entry(i)
                        matched_count += 1
                        if ds.get("preprocessed"):
                            already_processed = True
                    break

        msg = f"Successfully matched and updated {matched_count} out of {len(self.datasets)} datasets."
        if already_processed:
            msg += "\n\nNote: Some updated datasets were already processed. You will need to re-process them for the new geometry to take effect."
        
        QMessageBox.information(self, "Probendaten Import", msg)

    def remove_dataset(self):
        row = self.dataset_list.currentRow()
        if row < 0 or row >= self.dataset_list.count():
            return

        if row >= len(self.datasets):
            self.datasets = self.datasets[: self.dataset_list.count()]
            if row >= len(self.datasets):
                item = self.dataset_list.takeItem(row)
                del item
                self.current_index = None
                self.show_preview()
                self.refresh_ui_state()
                return

        reply = QMessageBox.question(
            self, "Remove Dataset",
            f"Really remove '{self.datasets[row]['filename']}'?",
            QMessageBox.Yes | QMessageBox.No
        )
        if reply != QMessageBox.Yes:
            return

        del self.datasets[row]
        item = self.dataset_list.takeItem(row)
        del item

        if len(self.datasets) == 0 or self.dataset_list.count() == 0:
            self.dataset_list.blockSignals(True)
            self.dataset_list.clear()
            self.dataset_list.setCurrentRow(-1)
            self.dataset_list.blockSignals(False)

            self.current_index = None
            self.show_preview()
            self.refresh_ui_state()
            return

        new_row = min(row, self.dataset_list.count() - 1)
        self.dataset_list.setCurrentRow(new_row)  
        self.refresh_ui_state()

    def clear_all_datasets(self):
        if not self.datasets:
            return

        reply = QMessageBox.question(
            self, "Clear All Datasets",
            "Are you sure you want to remove ALL loaded datasets? This cannot be undone.",
            QMessageBox.Yes | QMessageBox.No
        )
        
        if reply == QMessageBox.Yes:
            self.datasets.clear()
            self.current_index = None
            
            self.dataset_list.blockSignals(True)
            self.dataset_list.clear()
            self.dataset_list.blockSignals(False)
            
            self.show_preview()
            self.refresh_ui_state()
            self.update_buttons(-1)

    # -------------------- Processing --------------------
    def run_preprocess(self):
        if self.current_index is None or self.current_index < 0 or self.current_index >= len(self.datasets):
            QMessageBox.warning(self, "Error", "Select a dataset first.")
            return

        ds0 = self.datasets[self.current_index]

        dlg = PreprocessOptionsDialog(
            self,
            default_strain_increment=ds0.get("strain_increment", 0.005),
        )
        if dlg.exec() != QDialog.Accepted:
            return

        auto_average = dlg.get_auto_average() 
        dx = float(dlg.get_strain_increment())

        if self.chk_process_all.isChecked():
            targets = self.datasets
        else:
            targets = [self.datasets[self.current_index]]

        processed_count = 0

        for ds in targets:
            try:
                ds["strain_increment"] = dx

                ds["proc_data"] = preprocess_data(
                    ds["raw_data"],
                    ds["diameter"],
                    ds["height"],
                    strain_increment=dx,
                )
                ds["preprocessed"] = True
                processed_count += 1

                self.update_list_entry(self.datasets.index(ds))

            except Exception as e:
                QMessageBox.critical(self, "Preprocessing error", f"{ds['filename']}:\n{str(e)}")
                return

        self.refresh_ui_state()
        
        msg = f"Processed {processed_count} dataset(s)."

        if auto_average and processed_count > 0:
            avg_count = self._run_averaging()
            if avg_count > 0:
                msg += f"\n\nAutomatically generated {avg_count} averaged dataset(s) at the bottom of the list."

        if processed_count > 0:
            QMessageBox.information(self, "Success", msg)

    def _run_averaging(self):
        processed = [d for d in self.datasets if d.get("preprocessed", False) and d.get("proc_data") is not None]
        if not processed:
            return 0

        groups = {}
        for ds in processed:
            if str(ds["filename"]).startswith("AVG_"):
                continue

            raw_name = ds["filename"].replace(".csv", "")
            parts = raw_name.split("-")
            
            if len(parts) > 1:
                condition = "-".join(parts[1:])
            else:
                condition = raw_name 
                
            if condition not in groups:
                groups[condition] = []
            groups[condition].append(ds)

        new_datasets = []
        averaged_count = 0

        for condition, group in groups.items():
            if len(group) < 2:
                continue

            min_len = min(len(ds["proc_data"]["x"]) for ds in group)

            avg_df = pd.DataFrame()
            avg_df["x"] = group[0]["proc_data"]["x"].values[:min_len]

            cols_to_avg = ["y_none", "y_mild", "y_strong", "temperature"]
            for col in cols_to_avg:
                if col in group[0]["proc_data"].columns:
                    stacked = np.vstack([ds["proc_data"][col].values[:min_len] for ds in group])
                    avg_df[col] = np.mean(stacked, axis=0)

            avg_ds = {
                "filename": f"AVG_{condition}.csv",
                "raw_data": avg_df, 
                "proc_data": avg_df,
                "preprocessed": True,
                "diameter": group[0]["diameter"],
                "height": group[0]["height"],
                "strain_increment": group[0].get("strain_increment", 0.005),
            }
            new_datasets.append(avg_ds)
            averaged_count += 1

        for nds in new_datasets:
            existing_names = [d["filename"] for d in self.datasets]
            if nds["filename"] in existing_names:
                idx = existing_names.index(nds["filename"])
                self.datasets[idx] = nds
                self.update_list_entry(idx)
            else:
                self.datasets.append(nds)
                self.update_list_entry(len(self.datasets) - 1)

        self.refresh_ui_state()
        return averaged_count

    # -------------------- PLOTTING --------------------
    def _build_figure(self):
        processed = [
            d for d in self.datasets
            if d.get("preprocessed", False) and d.get("proc_data") is not None
        ]
        if not processed:
            QMessageBox.warning(self, "Error", "No processed datasets to plot.")
            return None

        selected = []
        if self.chk_none.isChecked():
            selected.append("none")
        if self.chk_mild.isChecked():
            selected.append("mild")
        if self.chk_strong.isChecked():
            selected.append("strong")

        if not selected:
            QMessageBox.warning(self, "Error", "Select at least one curve to plot.")
            return None

        FILTER_LINESTYLES = {
            "none": "-",
            "mild": "--",
            "strong": "-.",
        }

        num_items = len(processed)
        if num_items <= 15:
            leg_cols = 1
            plot_width = 0.75
            font_sz = 8
            fig_width = 9
        elif num_items <= 30:
            leg_cols = 2
            plot_width = 0.60
            font_sz = 7
            fig_width = 12
        else:
            leg_cols = 3
            plot_width = 0.45
            font_sz = 6
            fig_width = 15

        fig, axes = plt.subplots(2, 1, figsize=(fig_width, 10), sharex=False)
        plotted_any = False

        for ds in processed:
            df = ds["proc_data"]
            fname = ds["filename"]

            first_filt = next(
                (f for f in selected if f"y_{f}" in df.columns),
                None
            )
            if first_filt is None:
                continue

            line, = axes[0].plot(
                df["x"].values,
                df[f"y_{first_filt}"].values,
                linestyle=FILTER_LINESTYLES[first_filt],
                linewidth=2,
                label=fname,
            )
            dataset_color = line.get_color()
            plotted_any = True

            for filt in selected:
                if filt == first_filt:
                    continue
                col = f"y_{filt}"
                if col not in df.columns:
                    continue

                axes[0].plot(
                    df["x"].values,
                    df[col].values,
                    color=dataset_color,
                    linestyle=FILTER_LINESTYLES[filt],
                    linewidth=2 if filt != "none" else 1.2,
                    label=None,
                )

        axes[0].set_title("Flow Stress")
        axes[0].set_xlabel("Strain [-]")
        axes[0].set_ylabel("Flow Stress [MPa]")
        axes[0].grid(True)

        if not plotted_any:
            QMessageBox.warning(
                self,
                "Plotting error",
                "No flow-stress curves were plotted.\nCheck processed data columns."
            )
            plt.close(fig)
            return None

        has_temp = any("temperature" in d["proc_data"].columns for d in processed)
        if has_temp:
            for ds in processed:
                df = ds["proc_data"]
                if "temperature" not in df.columns:
                    continue

                tx = (
                    df["temperature_x"].values
                    if "temperature_x" in df.columns
                    else df["x"].values
                )

                axes[1].plot(tx, df["temperature"].values, label=ds["filename"])

            axes[1].set_title("Temperature (reference)")
            axes[1].set_xlabel("Strain [-]")
            axes[1].set_ylabel("Temperature [°C]")
            axes[1].set_ylim(0, 700)
            axes[1].grid(True)
        else:
            axes[1].set_visible(False)

        for ax in axes:
            if ax.get_visible():
                h, l = ax.get_legend_handles_labels()
                if h:
                    ax.legend(
                        h, l,
                        loc="upper left",
                        bbox_to_anchor=(1.02, 1.0),
                        borderaxespad=0.0,
                        fontsize=font_sz,
                        ncol=leg_cols
                    )

        plt.tight_layout(rect=[0, 0, plot_width, 1], h_pad=3.0)
        return fig

    def show_plots(self):
        fig = self._build_figure()
        if not fig:
            return

        # State tracking for the interactive cutter
        self._is_cutting = False
        self._hover_line = None
        
        # Check if ANY dataset has been processed (no longer restricted to one selected dataset)
        processed_datasets = [d for d in self.datasets if d.get("preprocessed") and d.get("proc_data") is not None]
        
        if processed_datasets:
            from matplotlib.widgets import Button
            
            # 1. Add the Matplotlib Button to the bottom right of the figure window
            self._btn_ax = fig.add_axes([0.78, 0.03, 0.18, 0.06])
            self._cut_btn = Button(self._btn_ax, 'Manual Offset\n(All Data)')
            
            main_ax = fig.axes[0]
            
            # 2. Button Click Event -> Activates Cutting Mode
            def on_click_btn(event):
                self._is_cutting = True
                main_ax.set_title("Hover & Click to set new origin for ALL curves", color='red', fontweight='bold')
                fig.canvas.draw_idle()
                
            self._cut_btn.on_clicked(on_click_btn)
            
            # 3. Mouse Hover Event -> Draws the vertical dashed line
            def on_mouse_move(event):
                if not self._is_cutting:
                    return
                # Ensure the mouse is actually inside the graph area
                if event.inaxes != main_ax:
                    if self._hover_line:
                        self._hover_line.set_visible(False)
                        fig.canvas.draw_idle()
                    return
                    
                # Update or create the red dashed cursor line
                if self._hover_line is None:
                    self._hover_line = main_ax.axvline(x=event.xdata, color='red', linestyle='--', linewidth=1.5)
                else:
                    self._hover_line.set_xdata([event.xdata, event.xdata])
                    self._hover_line.set_visible(True)
                fig.canvas.draw_idle()
                
            # 4. Plot Click Event -> Executes the Cut
            def on_plot_click(event):
                if not self._is_cutting:
                    return
                if event.inaxes != main_ax:
                    return
                
                # Capture the X strain coordinate clicked
                x_val = event.xdata
                
                # Reset interaction state
                self._is_cutting = False
                
                # Apply the mathematical cut to ALL datasets
                self._apply_manual_cut_all(x_val)
                
                # Close the window and use a timer to instantly reopen it with the new data
                import matplotlib.pyplot as plt
                from PySide6.QtCore import QTimer
                plt.close(fig)
                QTimer.singleShot(100, self.show_plots)

            # Bind the events to the figure canvas
            fig.canvas.mpl_connect('motion_notify_event', on_mouse_move)
            fig.canvas.mpl_connect('button_press_event', on_plot_click)

        import matplotlib.pyplot as plt
        plt.show()

    def _apply_manual_cut_all(self, target_x):
        """
        Slices ALL preprocessed dataframes at the user-selected point and shifts 
        the strain (X) back to 0. Stress (Y) remains strictly absolute.
        """
        import numpy as np
        
        for ds in self.datasets:
            df = ds.get("proc_data")
            
            # Skip if dataset hasn't been processed yet
            if df is None or "x" not in df.columns:
                continue
                
            x_vals = df["x"].values
            
            # Find the index of the data point closest to where the user clicked for THIS specific curve
            closest_idx = (np.abs(x_vals - target_x)).argmin()
            
            # Slice the dataset
            df_cut = df.iloc[closest_idx:].copy()
            
            if df_cut.empty:
                continue
            
            # Shift the strain axis so the new start point is strictly 0.0
            x0 = df_cut["x"].iloc[0]
            df_cut["x"] = df_cut["x"] - x0
            
            # If there's an independent temperature X-axis tracking, shift it too
            if "temperature_x" in df_cut.columns:
                df_cut["temperature_x"] = df_cut["temperature_x"] - x0
                
            # Overwrite the dataset state
            ds["proc_data"] = df_cut.reset_index(drop=True)
            
        # Refresh the PySide preview table so the GUI stays in sync
        self.show_preview()
    # -------------------- EXPORT LOGIC --------------------
    def _safe_sheet_name(self, name: str, suffix: str) -> str:
        bad = [":", "\\", "/", "?", "*", "[", "]"]
        out = name
        for b in bad:
            out = out.replace(b, "_")
        out = f"{out}__{suffix}"
        return out[:31]
    
    def export_results(self):
        if not self.datasets:
            QMessageBox.warning(self, "Error", "No datasets loaded.")
            return

        choice = QMessageBox.question(
            self,
            "Export scope",
            "Export all processed datasets?\n\n"
            "Yes  = Export ALL processed datasets (one sheet per dataset)\n"
            "No   = Export SELECTED dataset only",
            QMessageBox.Yes | QMessageBox.No | QMessageBox.Cancel,
            QMessageBox.No,  
        )
        if choice == QMessageBox.Cancel:
            return

        export_all = (choice == QMessageBox.Yes)

        if export_all:
            targets = [d for d in self.datasets if d.get("preprocessed", False) and d.get("proc_data") is not None]
            if not targets:
                QMessageBox.warning(self, "Error", "No processed datasets to export.")
                return
        else:
            if self.current_index is None or self.current_index < 0 or self.current_index >= len(self.datasets):
                QMessageBox.warning(self, "Error", "Select a dataset first.")
                return
            ds = self.datasets[self.current_index]
            if not ds.get("preprocessed", False) or ds.get("proc_data") is None:
                QMessageBox.warning(self, "Error", "Selected dataset is not processed yet.")
                return
            targets = [ds]

        export_filter = self.export_filter_combo.currentText().strip().lower()
        col_map = {"none": "y_none", "mild": "y_mild", "strong": "y_strong"}
        y_col = col_map.get(export_filter)
        if y_col is None:
            QMessageBox.warning(self, "Error", f"Unknown export filter: {export_filter}")
            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export Results",
            "",
            "Excel File (*.xlsx)"
        )
        if not file_path:
            return
        if not file_path.lower().endswith(".xlsx"):
            file_path += ".xlsx"

        try:
            with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
                for ds in targets:
                    df = ds["proc_data"]
                    if "x" not in df.columns or y_col not in df.columns:
                        raise ValueError(
                            f"{ds['filename']} missing required columns. Need 'x' and '{y_col}'."
                        )

                    out = pd.DataFrame({
                        "strain": df["x"].values,
                        "flow_stress": df[y_col].values,
                    })

                    if "temperature" in df.columns:
                        out["temperature"] = df["temperature"].values

                    sheet = self._safe_sheet_name(ds["filename"], f"filter_{export_filter}")
                    out.to_excel(writer, sheet_name=sheet, index=False)

                    # --- ALWAYS GENERATE EXCEL CHART ---
                    ws = writer.sheets[sheet]
                    chart = ScatterChart()
                    chart.title = f"Flow Curve: {ds['filename']}"
                    chart.style = 2  
                    
                    chart.x_axis.delete = False
                    chart.y_axis.delete = False
                    
                    chart.x_axis.tickLblPos = "low"
                    chart.y_axis.tickLblPos = "low"
                    
                    chart.x_axis.title = "Strain [-]"
                    chart.y_axis.title = "Flow Stress [MPa]"
                    chart.width = 16  
                    chart.height = 10 
                    chart.legend = None 

                    max_row = len(out) + 1
                    xvalues = Reference(ws, min_col=1, min_row=2, max_row=max_row)
                    values = Reference(ws, min_col=2, min_row=2, max_row=max_row)

                    series = Series(values, xvalues, title="Flow Stress")
                    chart.series.append(series)

                    ws.add_chart(chart, "E2")

            QMessageBox.information(
                self,
                "Success",
                f"Exported {len(targets)} dataset(s)\n"
                f"Scope: {'ALL processed' if export_all else 'Selected only'}\n"
                f"Filter: {export_filter}\n"
                f"File: {file_path}"
            )

        except Exception as e:
            QMessageBox.critical(self, "Error", f"Export failed:\n{str(e)}")

    def export_deform_qform(self):
        idx = self.current_index
        if idx is None or idx < 0 or idx >= len(self.datasets):
            QMessageBox.warning(self, "Error", "Select a dataset first.")
            return

        ds = self.datasets[idx]
        if not ds.get("preprocessed", False) or ds.get("proc_data") is None:
            QMessageBox.warning(self, "Error", "Process the dataset first.")
            return

        df = ds["proc_data"]

        try:
            filter_text = self.export_filter_combo.currentText().lower()
            if "mild" in filter_text:
                chosen_y_col = "y_mild"
            elif "strong" in filter_text:
                chosen_y_col = "y_strong"
            else:
                chosen_y_col = "y_none"

            out = deform_qform_table(df, y_col=chosen_y_col, n_points=25, cluster_factor=50.0)
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Failed to build DEFORM/QForm table:\n{str(e)}")
            return

        file_path, _ = QFileDialog.getSaveFileName(self, "Export DEFORM/QForm Table", "", "Excel File (*.xlsx)")
        if not file_path:
            return
        if not file_path.lower().endswith(".xlsx"):
            file_path += ".xlsx"

        try:
            with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
                sheet = self._safe_sheet_name(ds["filename"], "deform_qform")
                out.to_excel(writer, sheet_name=sheet, index=False)

                # --- ALWAYS GENERATE DEFORM EXCEL CHART ---
                ws = writer.sheets[sheet]
                chart = ScatterChart()
                chart.title = f"DEFORM/QForm Curve: {ds['filename']}"
                chart.style = 2  
                
                chart.x_axis.delete = False
                chart.y_axis.delete = False
                
                chart.x_axis.tickLblPos = "low"
                chart.y_axis.tickLblPos = "low"
                
                chart.x_axis.title = "Strain [-]"
                chart.y_axis.title = "Flow Stress [MPa]"
                chart.width = 16  
                chart.height = 10 
                chart.legend = None 

                max_row = len(out) + 1
                xvalues = Reference(ws, min_col=1, min_row=2, max_row=max_row)
                values = Reference(ws, min_col=2, min_row=2, max_row=max_row)

                series = Series(values, xvalues, title="Flow Stress")
                chart.series.append(series)

                ws.add_chart(chart, "D2") 

            QMessageBox.information(self, "Success", f"Exported {len(out)} points to:\n{file_path}")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Export failed:\n{str(e)}")

    def export_force_displacement_selected(self):
        idx = self.current_index
        if idx is None or idx < 0 or idx >= len(self.datasets):
            QMessageBox.warning(self, "Error", "Select a dataset first.")
            return

        ds = self.datasets[idx]
        if not ds.get("preprocessed", False) or ds.get("proc_data") is None:
            QMessageBox.warning(self, "Error", "Process the dataset first.")
            return

        raw = ds.get("raw_data", None)
        if raw is None or raw.empty:
            QMessageBox.warning(self, "Error", "No raw data available for this dataset.")
            return

        try:
            out = force_displacement_curve_from_raw(
                raw=raw,
                diameter=ds["diameter"],
                height=ds["height"],
                compliance=0.0029,
                failure_drop_frac=0.85,
            )
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Failed to build force–displacement curve:\n{str(e)}")
            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export Force–Displacement (Selected)",
            "",
            "Excel File (*.xlsx)"
        )
        if not file_path:
            return
        if not file_path.lower().endswith(".xlsx"):
            file_path += ".xlsx"

        try:
            with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
                sheet = self._safe_sheet_name(ds["filename"], "force_disp")
                out.to_excel(writer, sheet_name=sheet, index=False)

            QMessageBox.information(
                self,
                "Success",
                f"Exported force–displacement ({len(out)} rows)\nFile: {file_path}"
            )
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Export failed:\n{str(e)}")

    def run_extrapolation(self):
        idx = self.current_index
        if idx is None or idx < 0 or idx >= len(self.datasets):
            QMessageBox.warning(self, "Error", "Select a dataset first.")
            return

        ds = self.datasets[idx]
        if not ds.get("preprocessed", False) or ds.get("proc_data") is None:
            QMessageBox.warning(self, "Error", "Process the dataset first.")
            return

        # 1. Ask for Target Extrapolation Strain
        max_strain, ok = QInputDialog.getDouble(
            self, "Extrapolation Target", 
            "Enter maximum strain for extrapolation:", 
            2.0, 0.1, 100.0, 2
        )
        if not ok:
            return

        df = ds["proc_data"]
        
        # Determine which curve to fit
        filter_text = self.export_filter_combo.currentText().strip().lower()
        if "mild" in filter_text:
            y_col = "y_mild"
        elif "strong" in filter_text:
            y_col = "y_strong"
        else:
            y_col = "y_none"

        if y_col not in df.columns:
            QMessageBox.warning(self, "Error", f"Column {y_col} not found in processed data.")
            return

        epsilon = df["x"].values.astype(float)
        sigma = df[y_col].values.astype(float)
        
        # --- FIX: Only fit the uniform hardening region (up to peak stress) ---
        peak_idx = np.argmax(sigma)
        epsilon_fit = epsilon[:peak_idx + 1]
        sigma_fit = sigma[:peak_idx + 1]

        # 2. Fit Models
        from utils import fit_extrapolation_models, evaluate_model
        try:
            params_dict = fit_extrapolation_models(epsilon_fit, sigma_fit)
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Model fitting failed:\n{str(e)}")
            return
            
        # 3. Generate Extrapolation DataFrame
        interp_epsilon = np.linspace(max(0.01, np.min(epsilon)), max_strain, 200)
        out_data = {'epsilon': interp_epsilon}
        model_names = ['Hollomon', 'Ludwik', 'Swift', 'Voce', 'Hockett-Sherby', 'Swift+Voce']
        
        for model in model_names:
            safe_name = model.lower().replace('-', '_').replace('+', '_')
            out_data[safe_name] = evaluate_model(
                model, interp_epsilon, params_dict[model], 
                params_dict.get('Swift'), params_dict.get('Voce')
            )
            
        interp_df = pd.DataFrame(out_data)
            
        # 4. Visualizer
        import matplotlib.pyplot as plt
        fig, ax = plt.subplots(figsize=(10, 8))
        ax.scatter(epsilon, sigma, label="Measured Data", color="black", s=20, alpha=0.8)
        
        if params_dict['Hollomon'] is not None:
            ax.plot(interp_epsilon, out_data['hollomon'], label="Hollomon", linestyle="--")
            
        if params_dict['Ludwik'] is not None:
            ax.plot(interp_epsilon, out_data['ludwik'], label="Ludwik", linestyle="-.")
            
        if params_dict['Swift'] is not None:
            ax.plot(interp_epsilon, out_data['swift'], label="Swift", linestyle=":")
            
        if params_dict['Voce'] is not None:
            ax.plot(interp_epsilon, out_data['voce'], label="Voce", linestyle="-")
            
        if params_dict['Hockett-Sherby'] is not None:
            ax.plot(interp_epsilon, out_data['hockett_sherby'], label="Hockett-Sherby", linestyle=(0, (3, 1, 1, 1)))
            
        if params_dict['Swift+Voce'] is not None:
            ax.plot(interp_epsilon, out_data['swift_voce'], label="Swift+Voce", linewidth=2)
            
        ax.set_xlabel("Strain [-]")
        ax.set_ylabel("Flow Stress [MPa]")
        ax.set_title(f"Extrapolation Models: {ds['filename']}")
        ax.grid(True, alpha=0.3)
        
        # Legend at the bottom in 2 columns
        ax.legend(loc='upper center', bbox_to_anchor=(0.5, -0.12), ncol=2, borderaxespad=0.0)
        fig.subplots_adjust(bottom=0.25)
        
        plt.show()

       # 5. Optional Export
        reply = QMessageBox.question(
            self, "Export Results", 
            "Would you like to export the extrapolated curve data to Excel?", 
            QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes
        )

        if reply == QMessageBox.Yes:
            file_path, _ = QFileDialog.getSaveFileName(
                self,
                "Export Extrapolated Curves",
                self._safe_sheet_name(ds['filename'], 'extrapolated'),
                "Excel File (*.xlsx)"
            )
            
            if file_path:
                if not file_path.lower().endswith(".xlsx"):
                    file_path += ".xlsx"
                try:
                    from openpyxl.chart import ScatterChart, Reference, Series
                    
                    sheet_name = self._safe_sheet_name(ds['filename'], 'extrapolate')
                    
                    # Use ExcelWriter to write the data and embed the chart
                    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
                        interp_df.to_excel(writer, sheet_name=sheet_name, index=False)
                        
                        ws = writer.sheets[sheet_name]
                        chart = ScatterChart()
                        chart.title = f"Extrapolation Models: {ds['filename']}"
                        chart.style = 2  
                        chart.x_axis.title = "Strain [-]"
                        chart.y_axis.title = "Flow Stress [MPa]"
                        chart.width = 18  
                        chart.height = 12
                        
                        max_row = len(interp_df) + 1
                        # X-values are always the first column (epsilon)
                        xvalues = Reference(ws, min_col=1, min_row=2, max_row=max_row)
                        
                        # Loop through the remaining columns (the models) to add them to the chart
                        for col_idx, col_name in enumerate(interp_df.columns[1:], start=2):
                            # Only plot the model if it didn't fail (isn't full of NaNs)
                            if not interp_df[col_name].isna().all():
                                yvalues = Reference(ws, min_col=col_idx, min_row=2, max_row=max_row)
                                # Clean up the column name for the legend (e.g., 'hockett_sherby' -> 'Hockett Sherby')
                                nice_title = col_name.replace('_', ' ').title()
                                series = Series(yvalues, xvalues, title=nice_title)
                                chart.series.append(series)
                        
                        # Place the chart to the right of the data table (Column I)
                        ws.add_chart(chart, "I2")

                    QMessageBox.information(self, "Success", f"Extrapolated curves and graph exported to:\n{file_path}")
                except Exception as e:
                    QMessageBox.critical(self, "Error", f"Failed to save Excel file:\n{str(e)}")

    def export_deform_key(self):
        import re
        import numpy as np
        from scipy.interpolate import interp1d
        
        # 1. Gather all processed datasets
        processed_ds = [ds for ds in self.datasets if ds.get("preprocessed") and ds.get("proc_data") is not None]
        
        if not processed_ds:
            QMessageBox.warning(self, "Error", "No processed datasets available to export.")
            return

        # Determine which curve filter to use
        filter_text = self.export_filter_combo.currentText().strip().lower()
        y_col = "y_mild" if "mild" in filter_text else "y_strong" if "strong" in filter_text else "y_none"

        parsed_data = []
        max_strain_global = 100.0
        
        # 2. Parse Filenames for Temperature (T) and Strain Rate (SR)
        for ds in processed_ds:
            filename = ds['filename']
            match = re.search(r'T([\d\.]+).*?SR([\d\.]+)', filename, re.IGNORECASE)
            
            if not match:
                QMessageBox.warning(self, "Naming Error", 
                                    f"Could not extract Temperature and Strain Rate from:\n{filename}\n"
                                    "Ensure format contains e.g., 'T200' and 'SR1'.")
                return
                
            temp = float(match.group(1))
            sr = float(match.group(2))
            
            df = ds['proc_data']
            epsilon = df['x'].values
            sigma = df[y_col].values
            
            max_strain_global = min(max_strain_global, epsilon[-1])
            parsed_data.append({'T': temp, 'SR': sr, 'eps': epsilon, 'sig': sigma})

        # 3. Establish the 3D Grid Requirements
        unique_T = sorted(list(set(d['T'] for d in parsed_data)))
        unique_SR = sorted(list(set(d['SR'] for d in parsed_data)))
        
        if len(parsed_data) != len(unique_T) * len(unique_SR):
            QMessageBox.warning(self, "Grid Error", 
                                f"DEFORM requires a full matrix. You provided {len(parsed_data)} curves, "
                                f"but {len(unique_T)} Temps × {len(unique_SR)} Rates = {len(unique_T)*len(unique_SR)} curves are required.")
            return

        # Ask user for Material Name
        material_name, ok = QInputDialog.getText(self, "Material Name", "Enter Material Name for DEFORM:", text="Exported_Alloy")
        if not ok or not material_name: return

        # Ask user how many strain points they want
        num_strains, ok = QInputDialog.getInt(self, "Strain Points", "How many strain points for the matrix?", 20, 5, 100)
        if not ok: return

        common_strain = np.linspace(0.0, max_strain_global, num_strains)
        
        # 4. Generate the Matrix Data
        matrix_stresses = []
        for T in unique_T:
            for SR in unique_SR:
                curve = next(item for item in parsed_data if item['T'] == T and item['SR'] == SR)
                f = interp1d(curve['eps'], curve['sig'], bounds_error=False, fill_value=(curve['sig'][0], curve['sig'][-1]))
                interp_sig = np.maximum(f(common_strain), 0.01)
                matrix_stresses.extend(interp_sig.tolist())

        # Helper to format numbers like DEFORM expects (e.g., E+001 instead of E+01)
        def format_deform_float(val):
            s = f"{val:.10E}"
            base, exponent = s.split('E')
            return f"{base}E{exponent[0]}{int(exponent[1:]):03d}"

        def chunk_list(lst, chunk_size=7):
            lines = []
            for i in range(0, len(lst), chunk_size):
                chunk = lst[i:i+chunk_size]
                lines.append("    " + "    ".join([format_deform_float(v) for v in chunk]))
            return "\n".join(lines)

        file_path, _ = QFileDialog.getSaveFileName(self, "Save DEFORM Keyword File", f"{material_name}.key", "Keyword File (*.key);;Text File (*.txt)")
        if not file_path: return

        # 5. Format and Save the Keyword File
        static_footer = """YOUNG        2       0    2.1000000000E+005
POISON       2       0    3.0000000000E-001
EXPAND       2       0    1.2000000000E-005    2.0000000000E+001
THRCND       2       1      12
    1.0000000000E+002    5.0708000000E+001
    1.9900000000E+002    4.8112000000E+001
    2.9900000000E+002    4.5689000000E+001
    3.9900000000E+002    4.1718000000E+001
    4.9900000000E+002    3.8279000000E+001
    5.9900000000E+002    3.3943000000E+001
    6.9900000000E+002    3.0130000000E+001
    7.9900000000E+002    2.4747000000E+001
    9.9900000000E+002    3.2896000000E+001
    1.1990000000E+003    2.9756000000E+001
    1.3500000000E+003    2.9000000000E+001
    1.4850000000E+003    2.9000000000E+001
HEATCP       2       1      14       0
    1.0000000000E+002    3.8098100000E+000
    1.9900000000E+002    4.0397200000E+000
    2.4900000000E+002    4.1382000000E+000
    2.9900000000E+002    3.1254000000E+000
    3.4900000000E+002    4.4666800000E+000
    3.9900000000E+002    4.5980500000E+000
    4.9900000000E+002    5.0907000000E+000
    5.9900000000E+002    5.5505100000E+000
    6.9900000000E+002    6.0431500000E+000
    7.4900000000E+002    1.2414700000E+001
    7.9900000000E+002    4.8936400000E+000
    8.9900000000E+002    4.3024600000E+000
    1.3500000000E+003    4.3000000000E+000
    1.4850000000E+003    4.3000000000E+000
MASDEN       2       0    7.8700000000E-009
VSCOSY       2       0    0.0000000000E+000
COARSE       2       0
DIFBND       2       0    0.0000000000E+000
EMSVTY       2       0    7.0000000000E-001
HDNPHA       2       0    0.0000000000E+000
MSTMTR       2       0
CREEP        2       0
DIFCOE       2       0    0.0000000000E+000       0       1
RA1COF       2       0    0.0000000000E+000    0.0000000000E+000       1
RA2COF       2       0    0.0000000000E+000    0.0000000000E+000       1
ELRST        2       0    0.0000000000E+000
UTSDAT       2       0    0.0000000000E+000
HDNRUL       2       0
PMEAB        2       0    0.0000000000E+000
PMITT        2       0    0.0000000000E+000
MATDEN       2    0.0000000000E+000
BURGRS       2       0    0.0000000000E+000
ALPHA        2       0    0.0000000000E+000
NDISFM       2       0    0.0000000000E+000
RECVRY       2       0    0.0000000000E+000
SIZEMD       2       0
TXTURE       2       0       1       1
GBENGY       2       0    0.0000000000E+000
GBMOBI       2    0.0000000000E+000    0.0000000000E+000
NUCSIZ       2       0    0.0000000000E+000
HYPREL       2       0
PMCONS       2       0       0
"""

        try:
            with open(file_path, 'w') as f:
                f.write("*\n*  DEFORM MATERIAL KEYWORD FILE\n*\n")
                f.write("UNIT         1\n*\n*  Property Data of Material     2\n*\n")
                f.write(f"MTNAME       2\n{material_name}\n")
                f.write("FRAE2H       2    9.0000000000E-001       0\n")
                f.write("FPERV        2    0.0000000000E+000       0    0.0000000000E+000       0\n")
                f.write("FRCMOD       2       0    0.0000000000E+000    0.0000000000E+000    0.0000000000E+000    0.0000000000E+000    0.0000000000E+000       0       0\n")
                f.write("FSTRES       2       2\n")
                f.write(f"       {len(common_strain)}       {len(unique_SR)}       {len(unique_T)}\n")
                f.write(chunk_list(common_strain) + "\n")
                f.write(chunk_list(unique_SR) + "\n")
                f.write(chunk_list(unique_T) + "\n")
                f.write(chunk_list(matrix_stresses) + "\n")
                f.write(static_footer)
                
            QMessageBox.information(self, "Success", f"DEFORM 3D Keyword file successfully exported:\n{file_path}")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Failed to write file:\n{str(e)}")