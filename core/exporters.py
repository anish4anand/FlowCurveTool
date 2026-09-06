import re
import numpy as np
import pandas as pd
from scipy.interpolate import interp1d
from openpyxl.chart import ScatterChart, Reference, Series
from .constants import DEFORM_MATERIALS
from .material_models import deform_qform_table, force_displacement_curve_from_raw

def safe_sheet_name(name: str, suffix: str) -> str:
    bad = [":", "\\", "/", "?", "*", "[", "]"]
    out = name
    for b in bad: out = out.replace(b, "_")
    out = f"{out}__{suffix}"
    return out[:31]

def export_results_excel(processed_datasets: list, y_col: str, file_path: str):
    """Exports full processed flow curves to Excel with a scatter chart."""
    with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
        workbook = writer.book
        chart_sheet = workbook.create_sheet(title="Overview", index=0)
        chart = ScatterChart()
        chart.title = "True Stress vs True Strain"
        chart.style = 13
        chart.x_axis.title = 'True Strain'
        chart.y_axis.title = 'True Stress (MPa)'
        
        for ds in processed_datasets:
            sheet_name = str(ds.get('filename', 'Data'))[:31]
            df = ds['proc_data']
            
            export_df = pd.DataFrame({'True Strain': df['x'], 'True Stress (MPa)': df[y_col]})
            export_df.to_excel(writer, sheet_name=sheet_name, index=False)
            
            worksheet = writer.sheets[sheet_name]
            max_row = len(export_df) + 1
            xvalues = Reference(worksheet, min_col=1, min_row=2, max_row=max_row)
            yvalues = Reference(worksheet, min_col=2, min_row=2, max_row=max_row)
            
            series = Series(yvalues, xvalues, title_from_data=False)
            series.title = sheet_name
            chart.series.append(series)
            
        chart_sheet.add_chart(chart, "B2")

def export_deform_qform_table(processed_datasets: list, y_col: str, n_points: int, cluster_factor: float, file_path: str):
    """Applies geometric downsampling and saves to Excel with individual sheets and plots."""
    with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
        for ds in processed_datasets:
            filename = ds.get('filename', 'Unknown')
            sheet_name = safe_sheet_name(filename, "DEFORM")
            
            df = ds['proc_data']
            clustered_df = deform_qform_table(df, y_col=y_col, n_points=n_points, cluster_factor=cluster_factor)
            clustered_df.to_excel(writer, sheet_name=sheet_name, index=False)
            
            ws = writer.sheets[sheet_name]
            chart = ScatterChart()
            chart.title = f"DEFORM/QForm: {filename}"
            chart.style = 13  
            chart.x_axis.title = "True Strain [-]"
            chart.y_axis.title = "True Stress [MPa]"
            chart.width = 16
            chart.height = 10
            
            max_row = len(clustered_df) + 1
            xvalues = Reference(ws, min_col=1, min_row=2, max_row=max_row)
            yvalues = Reference(ws, min_col=2, min_row=2, max_row=max_row)
            
            series = Series(yvalues, xvalues, title_from_data=False, title="Flow Stress")
            chart.series.append(series)
            ws.add_chart(chart, "E2")

def export_force_displacement(datasets: list, diameter: float, height: float, file_path: str):
    """Recalculates and exports the force-displacement data to Excel with individual sheets."""
    valid_datasets = [ds for ds in datasets if ds.get('raw_data') is not None]
    if not valid_datasets:
        raise ValueError("No valid raw data available to export Force-Displacement.")
        
    with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
        for ds in valid_datasets:
            filename = ds.get('filename', 'Unknown')
            sheet_name = safe_sheet_name(filename, "FD")
            
            df_fd = force_displacement_curve_from_raw(ds['raw_data'], diameter, height)
            df_fd.to_excel(writer, sheet_name=sheet_name, index=False)
            
            ws = writer.sheets[sheet_name]
            chart = ScatterChart()
            chart.title = f"Force-Displacement: {filename}"
            chart.style = 13
            chart.x_axis.title = "Displacement [mm]"
            chart.y_axis.title = "Force [kN]"
            chart.width = 16
            chart.height = 10
            
            max_row = len(df_fd) + 1
            xvalues = Reference(ws, min_col=1, min_row=2, max_row=max_row)
            yvalues = Reference(ws, min_col=2, min_row=2, max_row=max_row)
            
            series = Series(yvalues, xvalues, title_from_data=False, title="Force vs Displacement")
            chart.series.append(series)
            ws.add_chart(chart, "E2")

def export_extrapolated_curves(file_path: str, sheet_name: str, interp_df: pd.DataFrame, ds_filename: str):
    """Exports evaluated extrapolation models to Excel with plotting."""
    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        interp_df.to_excel(writer, sheet_name=sheet_name, index=False)
        ws = writer.sheets[sheet_name]
        chart = ScatterChart()
        chart.title = f"Extrapolation Models: {ds_filename}"
        chart.style = 2  
        chart.x_axis.title = "Strain [-]"
        chart.y_axis.title = "Flow Stress [MPa]"
        chart.width = 18  
        chart.height = 12
        
        max_row = len(interp_df) + 1
        xvalues = Reference(ws, min_col=1, min_row=2, max_row=max_row)
        for col_idx, col_name in enumerate(interp_df.columns[1:], start=2):
            if not interp_df[col_name].isna().all():
                yvalues = Reference(ws, min_col=col_idx, min_row=2, max_row=max_row)
                nice_title = col_name.replace('_', ' ').title()
                series = Series(yvalues, xvalues, title=nice_title)
                chart.series.append(series)
        ws.add_chart(chart, "I2")

def format_deform_float(val: float) -> str:
    s = f"{val:.10E}"
    base, exponent = s.split('E')
    return f"{base}E{exponent[0]}{int(exponent[1:]):03d}"

def chunk_list(lst: list, chunk_size: int = 7) -> str:
    lines = []
    for i in range(0, len(lst), chunk_size):
        chunk = lst[i:i + chunk_size]
        lines.append("    " + "    ".join([format_deform_float(v) for v in chunk]))
    return "\n".join(lines)

def export_deform_keyword_file(processed_datasets: list, y_col: str, base_mat_key: str, custom_name: str, num_strains: int, file_path: str, cluster_factor: float = 50.0):
    """Interpolates the 3D matrix and writes the .key formatted string."""
    if base_mat_key not in DEFORM_MATERIALS:
        raise ValueError(f"Unknown material key: {base_mat_key}")

    mat_config = DEFORM_MATERIALS[base_mat_key]
    mat_id, static_footer = mat_config["id"], mat_config["footer"]
    parsed_data = []
    max_strain = 100.0

    for ds in processed_datasets:
        # Check if the UI already prompted and saved the metadata manually
        if 'T' in ds and 'SR' in ds:
            temp, sr = ds['T'], ds['SR']
        else:
            # Fallback to Regex with "RT" = 20°C logic
            filename = ds.get('filename', '')
            match = re.search(r'(?:T([\d\.]+)|(RT)).*?SR([\d\.]+)', filename, re.IGNORECASE)
            if not match:
                raise ValueError(f"Missing Temp/SR in filename: {filename}")
            
            temp = 20.0 if match.group(2) else float(match.group(1))
            sr = float(match.group(3))
        
        df = ds['proc_data']
        eps, sig = df['x'].values, df[y_col].values
        max_strain = min(max_strain, eps[-1])
        parsed_data.append({'T': temp, 'SR': sr, 'eps': eps, 'sig': sig})

    unique_T = sorted(list(set(d['T'] for d in parsed_data)))
    unique_SR = sorted(list(set(d['SR'] for d in parsed_data)))
    
    if len(parsed_data) != len(unique_T) * len(unique_SR):
        raise ValueError("Matrix incomplete. Missing Temp/Rate combinations.")

    t_array = np.linspace(0.0, 1.0, num_strains)
    common_strain = max_strain * (np.exp(cluster_factor * t_array / 10.0) - 1.0) / (np.exp(cluster_factor / 10.0) - 1.0)
    common_strain[0] = 0.0

    matrix_stresses = []
    for T in unique_T:
        for SR in unique_SR:
            curve = next(item for item in parsed_data if item['T'] == T and item['SR'] == SR)
            f = interp1d(curve['eps'], curve['sig'], bounds_error=False, fill_value=(curve['sig'][0], curve['sig'][-1]))
            matrix_stresses.extend(np.maximum(f(common_strain), 0.01).tolist())

    # Helper to format chunks of 5 values per line for DEFORM syntax
    def chunk_list(lst):
        return "\n".join(" ".join(f"{x:14.5E}" for x in lst[i:i+5]) for i in range(0, len(lst), 5))

    with open(file_path, 'w') as f:
        f.write("*\n*  DEFORM MATERIAL KEYWORD FILE\n*\n")
        f.write(f"UNIT         1\n*\n*  Property Data of Material     {mat_id}\n*\n")
        f.write(f"MTNAME       {mat_id}\n{custom_name}\n")
        f.write(f"FRAE2H       {mat_id}    9.0000000000E-001       0\n")
        f.write(f"FPERV        {mat_id}    0.0000000000E+000       0    0.0000000000E+000       0\n")
        f.write(f"FRCMOD       {mat_id}       0    0.0000000000E+000    0.0000000000E+000    0.0000000000E+000    0.0000000000E+000    0.0000000000E+000       0       0\n")
        f.write(f"FSTRES       {mat_id}       2\n")
        f.write(f"       {len(common_strain)}       {len(unique_SR)}       {len(unique_T)}\n")
        f.write(chunk_list(common_strain) + "\n")
        f.write(chunk_list(unique_SR) + "\n")
        f.write(chunk_list(unique_T) + "\n")
        f.write(chunk_list(matrix_stresses) + "\n")
        f.write(static_footer)